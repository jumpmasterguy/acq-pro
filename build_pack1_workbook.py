import copy, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

NAVY='FF0D1B2A'; TEAL='01696F'; TEALF='FF01696F'; LIGHT='FFE6F2F3'; GRAY='FF6B7280'
GREEN='FFC6EFCE'; YELLOW='FFFFEB9C'; RED='FFFFC7CE'
thin=Side(style='thin',color='FFE5E7EB'); B=Border(left=thin,right=thin,top=thin,bottom=thin)
RED_FONT=Font(color='FF9C0006',bold=True); GRN_FONT=Font(color='FF006100',bold=True); AMB_FONT=Font(color='FF9C5700',bold=True)

def hdr(ws,r,c,t,w=None):
    x=ws.cell(row=r,column=c,value=t); x.font=Font(bold=True,color='FFFFFFFF',size=10)
    x.fill=PatternFill('solid',fgColor=NAVY); x.border=B; x.alignment=Alignment(vertical='center',wrap_text=True)
    return x
def title(ws,r,t,sub=None):
    ws.cell(row=r,column=2,value=t).font=Font(bold=True,size=14,color=NAVY)
    if sub: ws.cell(row=r+1,column=2,value=sub).font=Font(size=10,italic=True,color=GRAY)
def lock(ws,unlocked):
    for rng in unlocked:
        for row in ws[rng]:
            for c in row: c.protection=Protection(locked=False)
    ws.protection.sheet=True; ws.protection.formatCells=False; ws.protection.formatColumns=False
    ws.protection.formatRows=False; ws.protection.selectLockedCells=False; ws.protection.selectUnlockedCells=False
def foot(ws,r):
    ws.cell(row=r,column=2,value='Acqlerate PM Essentials · acqlerate.com · September 2026 Edition · Reflects FAC 2025-06').font=Font(size=8,italic=True,color=GRAY)

# ───────── copy a sheet (values, styles, widths, DV, CF) ─────────
def copy_sheet(src,wb,name):
    t=wb.create_sheet(name)
    for row in src.iter_rows():
        for c in row:
            v=c.value
            n=t.cell(row=c.row,column=c.column,value=v)
            if c.has_style:
                n.font=copy.copy(c.font); n.fill=copy.copy(c.fill); n.border=copy.copy(c.border)
                n.alignment=copy.copy(c.alignment); n.number_format=c.number_format; n.protection=copy.copy(c.protection)
    for k,d in src.column_dimensions.items(): t.column_dimensions[k].width=d.width
    for k,d in src.row_dimensions.items():
        if d.height: t.row_dimensions[k].height=d.height
    for m in src.merged_cells.ranges: t.merge_cells(str(m))
    t.freeze_panes=src.freeze_panes
    for dv in src.data_validations.dataValidation:
        n=DataValidation(type=dv.type,formula1=dv.formula1,formula2=dv.formula2,allow_blank=dv.allow_blank,
                         showDropDown=dv.showDropDown,errorStyle=dv.errorStyle)
        n.error=dv.error; n.errorTitle=dv.errorTitle; t.add_data_validation(n)
        for r in str(dv.sqref).split(): n.add(r)
    for rng,rules in src.conditional_formatting._cf_rules.items():
        for rule in rules: t.conditional_formatting.add(str(rng.sqref),copy.copy(rule))
    t.protection=copy.copy(src.protection)
    t.page_setup.orientation=src.page_setup.orientation
    t.page_setup.fitToWidth=src.page_setup.fitToWidth; t.page_setup.fitToHeight=src.page_setup.fitToHeight
    t.sheet_properties.pageSetUpPr=copy.copy(src.sheet_properties.pageSetUpPr)
    if src.print_area: t.print_area=src.print_area
    return t

wb=openpyxl.Workbook(); wb.remove(wb.active)

# ══════════════════ START HERE ══════════════════
sh=wb.create_sheet('Start Here'); sh.column_dimensions['A'].width=2.5; sh.column_dimensions['B'].width=110
lines=[
('PM Essentials — Your First 90 Days as a DoD Program Manager',True,15),
('One workbook. Fill in the tabs, and the Dashboard builds your program status for you.',False,10),('',False,10),
('How this workbook fits together',True,12),
('1. Dashboard — reads every other tab and gives you a one-page program status. Print it for your PEO. You fill in only the header (program, PM, as-of date).',False,10),
('2. First 90 Days — the 30 things to do when you inherit a program, in order. Mark them off; the Dashboard tracks your progress.',False,10),
('3. Spend Plan — monthly obligation and expenditure plan vs. actuals for the fiscal year, against your appropriation\u2019s obligation goal. This is the tab you will live in.',False,10),
('4. Risk Register — 16 pre-loaded risks that actually kill DoD programs, scored and with mitigations. Rescore them for your program; the Risk Map redraws itself.',False,10),
('5. Risk Map — your risks plotted on the 5\u00d75 grid. Print this page for the weekly review.',False,10),
('6. RACI — 25 program decisions with who is Responsible, Accountable, Consulted, Informed. One row per decision, four fields, no grid to maintain.',False,10),
('7. PM Briefing Deck (separate PowerPoint) — 12 slides with the questions leadership will actually ask, in the speaker notes.',False,10),('',False,10),
('Where each tool is taught',True,12),
('Spend Plan and color of money \u2192 Module 2: Defense Finance & Budgeting, Lesson 1 and Lesson 8',False,10),
('Risk Register and Risk Map \u2192 Module 1: Foundations, Lesson 6 and Lesson 8',False,10),
('RACI and stakeholders \u2192 Module 6: Operations & Leadership, Lesson 5',False,10),
('Briefing Deck and the PMR \u2192 Module 6: Operations & Leadership, Lesson 8',False,10),
('Log in at acqlerate.com/app to open the lesson next to the tool.',False,10),('',False,10),
('Currency',True,12),
('Thresholds reflect FAC 2025-06 (SAT $350K, MPT $15K, effective Oct 1, 2025) and CMMC Phase 2 (C3PAO assessment required for Level 2, effective Nov 10, 2026). When these move, buyers get the updated file by email.',False,10),('',False,10),
('License & Use',True,12),
('Licensed for use on one program, by one organization. Customize text, branding, and layout freely. Company-wide deployment, sharing outside your organization, or resale requires an Acqlerate Team License (acqlerate.com/team).',False,10),
]
r=2
for t,b,s in lines:
    c=sh.cell(row=r,column=2,value=t); c.font=Font(bold=b,size=s,color=NAVY if b else '333333'); c.alignment=Alignment(wrap_text=True,vertical='top'); r+=1
sh.protection.sheet=True

# ══════════════════ RISK REGISTER + RISK MAP (copied) ══════════════════
rr=openpyxl.load_workbook('pack1_fixed_xlsx/risk-register.xlsx')
copy_sheet(rr['Risk Register'],wb,'Risk Register')
copy_sheet(rr['Risk Matrix Chart'],wb,'Risk Map')
# ══════════════════ RACI (copied) ══════════════════
rc=openpyxl.load_workbook('pack1_fixed_xlsx/stakeholder-raci.xlsx')
copy_sheet(rc['RACI Matrix'],wb,'RACI')
copy_sheet(rc['Instructions'],wb,'RACI Instructions')

# ══════════════════ FIRST 90 DAYS ══════════════════
f9=wb.create_sheet('First 90 Days')
title(f9,2,'First 90 Days — Inheriting a DoD Program','Thirty things to do, in order. Mark Status as you go; the Dashboard tracks each phase.')
for c,w in zip('ABCDEFG',[2.5,8,52,22,16,18,40]): f9.column_dimensions[c].width=w
for i,h in enumerate(['#','Action','Phase','Status','Owner','Notes / Evidence']): hdr(f9,5,2+i,h)
items=[
('Orient (Days 1\u201330)',[
 'Read the Acquisition Program Baseline and Acquisition Strategy. Know your cost, schedule, and performance thresholds before anyone asks.',
 'Schedule 1:1s in week one with your PEO, Contracting Officer, COR, and the contractor PM. Ask each: what worries you most?',
 'Pull the last three PMR decks. Read the risk and cost slides first \u2014 that is where the truth lives.',
 'Find the current spend plan and obligation status. Know your color of money and when each appropriation expires.',
 'Get read-in on every active contract: type, period of performance, CLIN structure, remaining options, ceiling.',
 'Review the contractor\u2019s last two CPARS and every open CAR or deficiency report.',
 'Confirm the RACI (RACI tab): who is Accountable for each decision. Fix the blanks now, not in a crisis.',
 'List every milestone, decision, and test event inside the next 90 days. These are your forcing functions.',
 'Verify CMMC / CUI posture on every contract handling CUI. Phase 2 (Nov 10, 2026) requires C3PAO assessment for Level 2.',
 'Stand up the Risk Register (Risk Register tab) with the outgoing PM\u2019s top ten. Do not rescore yet \u2014 just capture.']),
('Assess (Days 31\u201360)',[
 'Rescore every risk with your own probability and impact. Brief the Risk Map to your PEO within the phase.',
 'Reconcile obligations against plan (Spend Plan tab). Flag any appropriation at risk of expiring unobligated.',
 'Walk the Integrated Master Schedule with the contractor. Validate the critical path and where the float actually is.',
 'If EVM applies (>$20M cost-type), review CPI and SPI trends and the contractor\u2019s VAC. Ask what changed since last CPR.',
 'Audit CDRL delivery status. Which deliverables are late, and does the government owe review comments?',
 'Meet your DCMA and DCAA counterparts if assigned. Ask what they see in the contractor\u2019s systems.',
 'Assess staffing: gaps on the government team, and whether the contractor\u2019s key personnel are the ones proposed.',
 'Review every ECP and contract modification in the queue. Is each funded, and by which appropriation?',
 'Confirm test events in the next six months and whether entrance criteria are defined and owned.',
 'Draft your first PMR in your own format (Briefing Deck). Lead with risk and money.']),
('Act (Days 61\u201390)',[
 'Deliver your first PMR. Every red light gets a sentence on cause and a sentence on the fix.',
 'Publish the program battle rhythm: IPR cadence, risk board, spend reviews, and who attends each.',
 'Close or re-scope stale risks. Open new ones from your assessment. The register should look different than day 1.',
 'Submit spend plan adjustments to your resource manager before the next obligation review.',
 'Set interim CPARS expectations with the contractor in writing. No surprises at the annual rating.',
 'Lock the next milestone\u2019s entrance criteria and assign an owner to each one in the RACI.',
 'Start a decision log. Every decision in your first 90 days, dated, with who made it.',
 'Brief lessons learned to your PEO. Ask directly: what do you want done differently?',
 'Update the RACI for any role changes since day one.',
 'Schedule the six-month look-ahead review and put the date on the Dashboard.']),
]
r=6; n=1; phase_rows={}
for phase,acts in items:
    start=r
    for a in acts:
        f9.cell(row=r,column=2,value=n).border=B
        c=f9.cell(row=r,column=3,value=a); c.border=B; c.alignment=Alignment(wrap_text=True,vertical='top')
        f9.cell(row=r,column=4,value=phase).border=B
        s=f9.cell(row=r,column=5,value='Not Started'); s.border=B; s.alignment=Alignment(horizontal='center')
        f9.cell(row=r,column=6).border=B; f9.cell(row=r,column=7).border=B
        f9.row_dimensions[r].height=42
        if r%2==0:
            for cc in range(2,8): f9.cell(row=r,column=cc).fill=PatternFill('solid',fgColor='FFF9FAFB')
        r+=1; n+=1
    phase_rows[phase]=(start,r-1)
last=r-1
dv=DataValidation(type='list',formula1='"Not Started,In Progress,Done,N/A"',allow_blank=True); f9.add_data_validation(dv); dv.add(f'E6:E{last}')
f9.conditional_formatting.add(f'E6:E{last}',FormulaRule(formula=['E6="Done"'],fill=PatternFill('solid',fgColor=GREEN),font=GRN_FONT))
f9.conditional_formatting.add(f'E6:E{last}',FormulaRule(formula=['E6="In Progress"'],fill=PatternFill('solid',fgColor=YELLOW),font=AMB_FONT))
f9.freeze_panes='B6'
# progress block
pr=last+2
f9.cell(row=pr,column=2,value='PROGRESS').font=Font(bold=True,size=11,color=NAVY)
for i,h in enumerate(['Phase','Done','Total','% Complete']): hdr(f9,pr+1,3+i,h)
for i,(phase,_) in enumerate(items):
    rr_=pr+2+i; a,b_=phase_rows[phase]
    f9.cell(row=rr_,column=3,value=phase).border=B
    f9.cell(row=rr_,column=4,value=f'=COUNTIF(E{a}:E{b_},"Done")').border=B
    f9.cell(row=rr_,column=5,value=f'=COUNTA(E{a}:E{b_})-COUNTIF(E{a}:E{b_},"N/A")').border=B
    p=f9.cell(row=rr_,column=6,value=f'=IF(E{rr_}=0,0,D{rr_}/E{rr_})'); p.number_format='0%'; p.border=B
tr=pr+2+len(items)
f9.cell(row=tr,column=3,value='OVERALL').font=Font(bold=True)
f9.cell(row=tr,column=4,value=f'=SUM(D{pr+2}:D{tr-1})').font=Font(bold=True)
f9.cell(row=tr,column=5,value=f'=SUM(E{pr+2}:E{tr-1})').font=Font(bold=True)
p=f9.cell(row=tr,column=6,value=f'=IF(E{tr}=0,0,D{tr}/E{tr})'); p.number_format='0%'; p.font=Font(bold=True,color=TEAL)
F9_OVERALL=f'F{tr}'; F9_PHASE={ph:f'F{pr+2+i}' for i,(ph,_) in enumerate(items)}
foot(f9,tr+2)
lock(f9,[f'C6:G{last}'])
f9.page_setup.orientation='landscape'; f9.page_setup.fitToWidth=1; f9.page_setup.fitToHeight=0; f9.sheet_properties.pageSetUpPr.fitToPage=True

# ══════════════════ SPEND PLAN ══════════════════
sp=wb.create_sheet('Spend Plan')
title(sp,2,'Spend Plan — Obligations & Expenditures vs. Plan','One fiscal year, one appropriation line. Enter the plan, then post actuals each month. Status and the Dashboard update themselves.')
for c,w in zip('ABCDEFGHIJKLM',[2.5,14,16,16,16,16,15,13,13,16,16,16,16]): sp.column_dimensions[c].width=w
# header inputs
inp=[('Fiscal Year','FY2027'),('Appropriation','O&M'),('Total Budget Authority',4150000),('Years Available',None),('Obligation Goal (Yr 1)',None),('Goal $ (Yr 1)',None)]
for i,(k,v) in enumerate(inp):
    sp.cell(row=5+i,column=2,value=k).font=Font(bold=True,size=10)
    c=sp.cell(row=5+i,column=3,value=v); c.border=B
sp['C7'].number_format='$#,##0'
sp['C8']='=VLOOKUP(C6,$I$5:$K$8,2,FALSE)'
sp['C9']='=VLOOKUP(C6,$I$5:$K$8,3,FALSE)'; sp['C9'].number_format='0%'
sp['C10']='=C7*C9'; sp['C10'].number_format='$#,##0'
dvA=DataValidation(type='list',formula1='"O&M,RDT&E,Procurement"',allow_blank=False); sp.add_data_validation(dvA); dvA.add('C6')
# goal table
for i,h in enumerate(['Appropriation','Yrs Available','Yr-1 Oblig Goal']): hdr(sp,4,9+i,h)
for i,(a,y,g) in enumerate([('O&M',1,1.00),('RDT&E',2,0.90),('Procurement',3,0.80)]):
    sp.cell(row=5+i,column=9,value=a).border=B; sp.cell(row=5+i,column=10,value=y).border=B
    c=sp.cell(row=5+i,column=11,value=g); c.number_format='0%'; c.border=B
sp.cell(row=8,column=9,value='Typical OSD first-year obligation goals. Verify against your Component\u2019s current guidance \u2014 they change.').font=Font(size=8,italic=True,color=GRAY)
# monthly table
H=12
heads=['Month','Planned Oblig','Actual Oblig','Cum Planned','Cum Actual','Variance $','Variance %','Status','Planned Exp','Actual Exp','Cum Plan Exp','Cum Act Exp']
for i,h in enumerate(heads): hdr(sp,H,2+i,h)
months=['Oct','Nov','Dec','Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep']
# realistic S-curve plan on $4.15M; actuals posted through March
plan=[0.06,0.08,0.10,0.09,0.09,0.09,0.09,0.09,0.09,0.08,0.07,0.07]
actual=[0.05,0.07,0.11,0.08,0.10,0.08,None,None,None,None,None,None]
exp_plan=[0.02,0.04,0.06,0.08,0.09,0.10,0.10,0.10,0.10,0.10,0.11,0.10]
exp_act=[0.02,0.03,0.05,0.07,0.09,0.09,None,None,None,None,None,None]
BA=4150000
for i,m in enumerate(months):
    r=H+1+i
    sp.cell(row=r,column=2,value=m).border=B
    for col,val in [(3,plan[i]),(4,actual[i]),(10,exp_plan[i]),(11,exp_act[i])]:
        c=sp.cell(row=r,column=col,value=None if val is None else round(BA*val)); c.number_format='$#,##0'; c.border=B
    prev=f'E{r-1}' if i else '0'; prevA=f'F{r-1}' if i else '0'
    sp.cell(row=r,column=5,value=f'=C{r}+{prev}').number_format='$#,##0'
    sp.cell(row=r,column=6,value=f'=IF(D{r}="","",D{r}+N({prevA}))').number_format='$#,##0'
    sp.cell(row=r,column=7,value=f'=IF(F{r}="","",F{r}-E{r})').number_format='$#,##0;[Red]-$#,##0'
    sp.cell(row=r,column=8,value=f'=IF(F{r}="","",IF(E{r}=0,0,(F{r}-E{r})/E{r}))').number_format='0%;[Red]-0%'
    st=sp.cell(row=r,column=9,value=f'=IF(F{r}="","",IF(H{r}<-0.1,"BEHIND",IF(H{r}>0.1,"AHEAD","ON PLAN")))'); st.alignment=Alignment(horizontal='center'); st.font=Font(bold=True,size=9)
    pe=f'L{r-1}' if i else '0'; pa=f'M{r-1}' if i else '0'
    sp.cell(row=r,column=12,value=f'=J{r}+{pe}').number_format='$#,##0'
    sp.cell(row=r,column=13,value=f'=IF(K{r}="","",K{r}+N({pa}))').number_format='$#,##0'
    for cc in range(5,14): sp.cell(row=r,column=cc).border=B
last=H+12
sp.conditional_formatting.add(f'I{H+1}:I{last}',FormulaRule(formula=[f'I{H+1}="BEHIND"'],fill=PatternFill('solid',fgColor=RED),font=RED_FONT))
sp.conditional_formatting.add(f'I{H+1}:I{last}',FormulaRule(formula=[f'I{H+1}="AHEAD"'],fill=PatternFill('solid',fgColor=YELLOW),font=AMB_FONT))
sp.conditional_formatting.add(f'I{H+1}:I{last}',FormulaRule(formula=[f'I{H+1}="ON PLAN"'],fill=PatternFill('solid',fgColor=GREEN),font=GRN_FONT))
# summary
s=last+2
sp.cell(row=s,column=2,value='YEAR-TO-DATE').font=Font(bold=True,size=11,color=NAVY)
summ=[('Months posted',f'=COUNT(D{H+1}:D{last})','0'),
      ('Planned to date',f'=IF(C{s+1}=0,0,INDEX(E{H+1}:E{last},C{s+1}))','$#,##0'),
      ('Obligated to date',f'=IF(C{s+1}=0,0,INDEX(F{H+1}:F{last},C{s+1}))','$#,##0'),
      ('% of BA obligated',f'=IF(C7=0,0,C{s+3}/C7)','0%'),
      ('% of plan achieved',f'=IF(C{s+2}=0,0,C{s+3}/C{s+2})','0%'),
      ('Remaining to Yr-1 goal',f'=MAX(0,C10-C{s+3})','$#,##0'),
      ('Expended to date',f'=IF(C{s+1}=0,0,INDEX(M{H+1}:M{last},C{s+1}))','$#,##0'),
      ('Overall status',f'=IF(C{s+1}=0,"No actuals posted",IF(C{s+5}<0.9,"BEHIND PLAN",IF(C{s+5}>1.1,"AHEAD OF PLAN","ON PLAN")))','@')]
for i,(k,f,nf) in enumerate(summ):
    sp.cell(row=s+1+i,column=2,value=k).font=Font(bold=True,size=10)
    c=sp.cell(row=s+1+i,column=3,value=f); c.number_format=nf; c.border=B
SP_STATUS=f'C{s+8}'; SP_PCT_BA=f'C{s+4}'; SP_PCT_PLAN=f'C{s+5}'; SP_OBL=f'C{s+3}'; SP_PLAN=f'C{s+2}'; SP_REM=f'C{s+6}'
sp.conditional_formatting.add(SP_STATUS,FormulaRule(formula=[f'LEFT({SP_STATUS},6)="BEHIND"'],fill=PatternFill('solid',fgColor=RED),font=RED_FONT))
sp.conditional_formatting.add(SP_STATUS,FormulaRule(formula=[f'{SP_STATUS}="ON PLAN"'],fill=PatternFill('solid',fgColor=GREEN),font=GRN_FONT))
sp.cell(row=s+10,column=2,value='BEHIND / AHEAD trip at \u00b110% of cumulative plan. An expiring appropriation that is BEHIND in Q3 is the number-one avoidable failure in program finance \u2014 act in Q2.').font=Font(size=9,italic=True,color=GRAY)
foot(sp,s+12)
sp.freeze_panes=f'C{H+1}'
lock(sp,['C5:C7',f'C{H+1}:D{last}',f'J{H+1}:K{last}'])
sp.page_setup.orientation='landscape'; sp.page_setup.fitToWidth=1; sp.page_setup.fitToHeight=0; sp.sheet_properties.pageSetUpPr.fitToPage=True

# ══════════════════ DASHBOARD ══════════════════
db=wb.create_sheet('Dashboard',1)
for c,w in zip('ABCDEFGH',[2.5,26,22,4,26,22,4,30]): db.column_dimensions[c].width=w
db['B2']='PROGRAM DASHBOARD'; db['B2'].font=Font(bold=True,size=16,color=NAVY)
db['B3']='Builds itself from the other tabs. Fill in the three header cells, then print this page for your PEO.'; db['B3'].font=Font(size=10,italic=True,color=GRAY)
for i,(k,v) in enumerate([('Program','[Program Name]'),('Program Manager','[Your Name]'),('As of',datetime.date(2026,9,15))]):
    db.cell(row=5+i,column=2,value=k).font=Font(bold=True,size=10)
    c=db.cell(row=5+i,column=3,value=v); c.border=B; c.fill=PatternFill('solid',fgColor='FFFFFBEB')
db['C7'].number_format='mmm d, yyyy'
def box(r,c,t): 
    x=db.cell(row=r,column=c,value=t); x.font=Font(bold=True,size=11,color='FFFFFFFF'); x.fill=PatternFill('solid',fgColor=TEALF)
    db.cell(row=r,column=c+1).fill=PatternFill('solid',fgColor=TEALF)
def kv(r,c,k,f,nf='@',bold=False):
    db.cell(row=r,column=c,value=k).font=Font(size=10,bold=bold)
    x=db.cell(row=r,column=c+1,value=f); x.number_format=nf; x.border=B; x.font=Font(size=10,bold=bold); x.alignment=Alignment(horizontal='right')
    return x
RR="'Risk Register'"; SPN="'Spend Plan'"; F9N="'First 90 Days'"; RA="'RACI'"
# money
box(9,2,'MONEY'); 
kv(10,2,'Appropriation',f"={SPN}!C6"); kv(11,2,'Budget Authority',f"={SPN}!C7",'$#,##0')
kv(12,2,'Obligated to date',f"={SPN}!{SP_OBL}",'$#,##0'); kv(13,2,'% of BA obligated',f"={SPN}!{SP_PCT_BA}",'0%')
kv(14,2,'% of plan achieved',f"={SPN}!{SP_PCT_PLAN}",'0%'); kv(15,2,'Remaining to Yr-1 goal',f"={SPN}!{SP_REM}",'$#,##0')
ms=kv(16,2,'Status',f"={SPN}!{SP_STATUS}",bold=True); ms.alignment=Alignment(horizontal='center')
# risk
box(9,5,'RISK')
kv(10,5,'HIGH',f'=COUNTIF({RR}!I16:I31,"HIGH")','0'); kv(11,5,'MEDIUM',f'=COUNTIF({RR}!I16:I31,"MEDIUM")','0'); kv(12,5,'LOW',f'=COUNTIF({RR}!I16:I31,"LOW")','0')
kv(13,5,'Open',f'=COUNTIF({RR}!N16:N31,"Open")','0')
rs=kv(14,5,'Status',f'=IF(F10>=3,"RED",IF(F10>=1,"AMBER","GREEN"))',bold=True); rs.alignment=Alignment(horizontal='center')
db.cell(row=15,column=5,value='HIGH risks (open)').font=Font(size=10,bold=True)
hl=db.cell(row=16,column=5)
hl.value=ArrayFormula(ref='E16',text=f'=_xlfn.TEXTJOIN(CHAR(10),TRUE,IF(({RR}!$I$16:$I$31="HIGH")*({RR}!$N$16:$N$31="Open"),{RR}!$B$16:$B$31&" "&{RR}!$C$16:$C$31,""))')
db.merge_cells('E16:F19'); hl.alignment=Alignment(wrap_text=True,vertical='top'); hl.font=Font(size=9); hl.border=B
# 90 days
box(21,2,'FIRST 90 DAYS')
kv(22,2,'Orient (1\u201330)',f"={F9N}!{F9_PHASE['Orient (Days 1\u201330)']}",'0%'); kv(23,2,'Assess (31\u201360)',f"={F9N}!{F9_PHASE['Assess (Days 31\u201360)']}",'0%')
kv(24,2,'Act (61\u201390)',f"={F9N}!{F9_PHASE['Act (Days 61\u201390)']}",'0%'); kv(25,2,'Overall',f"={F9N}!{F9_OVERALL}",'0%',bold=True)
# raci
box(21,5,'ACCOUNTABILITY')
kv(22,5,'Decisions tracked',f'=COUNTA({RA}!B8:B32)','0'); kv(23,5,'Missing an Accountable',f'=COUNTBLANK({RA}!D8:D32)','0')
as_=kv(24,5,'Status',f'=IF(F23=0,"GREEN","RED")',bold=True); as_.alignment=Alignment(horizontal='center')
# lights
for cell in ['C16','F14','F24']:
    db.conditional_formatting.add(cell,FormulaRule(formula=[f'OR({cell}="RED",LEFT({cell},6)="BEHIND")'],fill=PatternFill('solid',fgColor=RED),font=RED_FONT))
    db.conditional_formatting.add(cell,FormulaRule(formula=[f'OR({cell}="AMBER",LEFT({cell},5)="AHEAD")'],fill=PatternFill('solid',fgColor=YELLOW),font=AMB_FONT))
    db.conditional_formatting.add(cell,FormulaRule(formula=[f'OR({cell}="GREEN",{cell}="ON PLAN")'],fill=PatternFill('solid',fgColor=GREEN),font=GRN_FONT))
# how to read
db['H9']='HOW TO READ THIS'; db['H9'].font=Font(bold=True,size=11,color=NAVY)
notes=['MONEY goes red when cumulative obligations fall more than 10% behind plan. If the appropriation expires this FY, that is the slide to lead with.',
'RISK is red at three or more HIGH risks. The list shows every open HIGH \u2014 each one needs a mitigation owner named on the register.',
'FIRST 90 DAYS shows where you are in taking over the program. Below 50% on Orient after day 30 means you are still finding out things you should already know.',
'ACCOUNTABILITY is red if any decision in the RACI has no Accountable owner. That is the gap that turns into \u201cI thought you had that.\u201d']
for i,t in enumerate(notes):
    c=db.cell(row=10+i*3,column=8,value=t); c.font=Font(size=9,color='333333'); c.alignment=Alignment(wrap_text=True,vertical='top'); db.merge_cells(start_row=10+i*3,start_column=8,end_row=12+i*3,end_column=8)
foot(db,28)
lock(db,['C5:C7'])
db.page_setup.orientation='landscape'; db.page_setup.fitToWidth=1; db.page_setup.fitToHeight=1; db.sheet_properties.pageSetUpPr.fitToPage=True; db.print_area='A1:H28'

# order
order=['Start Here','Dashboard','First 90 Days','Spend Plan','Risk Register','Risk Map','RACI','RACI Instructions']
wb._sheets=[wb[n] for n in order]
wb.save('pm-essentials-workbook.xlsx')
print('built:',order)
