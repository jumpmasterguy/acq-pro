"""
FILE 4: wrap-rate-breakdown.xlsx
Wrap rate calculator, comparison, and instructions: 3 sheets.

Run from anywhere:  python3 client/public/products/pack3-finance-cheat-sheets/build_wrap_rate.py
Writes the workbook to BOTH twin folders (client/public/products/... and products/...).
Then recalculate so formula cells carry cached values for previewers.

Brand header uses brand/acqlerate-lockup-light.png (icon + Acq/lerate wordmark),
rendered by scripts/pack3/make_lockup_png.py. Never type the brand as text in caps.
"""
from pathlib import Path
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, Protection
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.drawing.image import Image as XLImage

HERE = Path(__file__).resolve().parent
ROOT = next(p for p in HERE.parents if (p / "brand" / "acqlerate-lockup-light.png").exists())
LOCKUP = ROOT / "brand" / "acqlerate-lockup-light.png"
EDITION = "September 2026 Edition"
FOOTER = f"Acqlerate Finance Cheat Sheets  ·  acqlerate.com  ·  {EDITION}  ·  Free to share unaltered. Educational reference, not official DoD guidance."

wb = openpyxl.Workbook()

# ── Brand colors ──
TEAL = "01696F"
NAVY = "0D1B2A"
WHITE = "FFFFFF"
LIGHT_TEAL = "E0F2F3"
VERY_LIGHT_TEAL = "F0FAFA"
DARK_GRAY = "333333"
MID_GRAY = "666666"
BORDER_GRAY = "C0C0C0"
LIGHT_GRAY = "F5F5F5"
INPUT_BLUE = "0000FF"  # Financial model convention
LIGHT_YELLOW = "FFFFF0"

# ── Fills ──
fill_navy = PatternFill("solid", fgColor=NAVY)
fill_teal = PatternFill("solid", fgColor=TEAL)
fill_light_teal = PatternFill("solid", fgColor=LIGHT_TEAL)
fill_white = PatternFill("solid", fgColor=WHITE)
fill_light_gray = PatternFill("solid", fgColor=LIGHT_GRAY)
fill_very_light_teal = PatternFill("solid", fgColor=VERY_LIGHT_TEAL)
fill_input = PatternFill("solid", fgColor="FFFFF0")  # Light yellow for inputs
fill_output = PatternFill("solid", fgColor="E8F5E9")  # Light green for outputs
fill_result = PatternFill("solid", fgColor=LIGHT_TEAL)

# ── Fonts ──
font_title = Font(name="Calibri", size=16, bold=True, color=WHITE)
font_subtitle = Font(name="Calibri", size=9, color="B0C4DE")
font_section = Font(name="Calibri", size=11, bold=True, color=WHITE)
font_header = Font(name="Calibri", size=9, bold=True, color=WHITE)
font_body = Font(name="Calibri", size=10, color=DARK_GRAY)
font_body_bold = Font(name="Calibri", size=10, bold=True, color=DARK_GRAY)
font_body_sm = Font(name="Calibri", size=9, color=DARK_GRAY)
font_input = Font(name="Calibri", size=11, bold=True, color=INPUT_BLUE)  # Blue = input
font_formula = Font(name="Calibri", size=11, color=DARK_GRAY)  # Black = formula
font_formula_bold = Font(name="Calibri", size=11, bold=True, color=DARK_GRAY)
font_result = Font(name="Calibri", size=13, bold=True, color=TEAL)
font_label = Font(name="Calibri", size=10, color=MID_GRAY)
font_label_bold = Font(name="Calibri", size=10, bold=True, color=TEAL)
font_note = Font(name="Calibri", size=8, italic=True, color=MID_GRAY)
font_footer = Font(name="Calibri", size=7, italic=True, color=MID_GRAY)
font_inst_body = Font(name="Calibri", size=10, color=DARK_GRAY)
font_inst_bold = Font(name="Calibri", size=10, bold=True, color=DARK_GRAY)
font_inst_head = Font(name="Calibri", size=11, bold=True, color=TEAL)

# ── Borders ──
thin = Border(
    left=Side(style="thin", color=BORDER_GRAY),
    right=Side(style="thin", color=BORDER_GRAY),
    top=Side(style="thin", color=BORDER_GRAY),
    bottom=Side(style="thin", color=BORDER_GRAY),
)
bottom_thin = Border(bottom=Side(style="thin", color=BORDER_GRAY))
bottom_thick = Border(bottom=Side(style="medium", color=TEAL))
input_border = Border(
    left=Side(style="thin", color=INPUT_BLUE),
    right=Side(style="thin", color=INPUT_BLUE),
    top=Side(style="thin", color=INPUT_BLUE),
    bottom=Side(style="thin", color=INPUT_BLUE),
)

# ── Alignments ──
ac = Alignment(horizontal="center", vertical="center", wrap_text=True)
al = Alignment(horizontal="left", vertical="center", wrap_text=True)
al_indent = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
ar = Alignment(horizontal="right", vertical="center")
al_top = Alignment(horizontal="left", vertical="top", wrap_text=True)

def apply_fill(ws, row, c1, c2, fill):
    for c in range(c1, c2+1):
        ws.cell(row=row, column=c).fill = fill


def brand_header(ws, title, subtitle, last_col, note=None):
    """Light header matching the newer packs: lockup, navy title, grey subtitle,
    thin teal rule. Content starts on row 6."""
    img = XLImage(str(LOCKUP))
    img.height = 26
    img.width = round(26 * 538 / 128)
    ws.add_image(img, "B1")
    ws.row_dimensions[1].height = 24
    last = get_column_letter(last_col)
    ws.merge_cells(f"B2:{last}2")
    ws["B2"].value = title
    ws["B2"].font = Font(name="Calibri", size=16, bold=True, color="0F172A")
    ws["B2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 26
    if note:
        prev = get_column_letter(last_col - 1)
        ws.merge_cells(f"B3:{prev}3")
        ws.cell(row=3, column=last_col).value = note
        ws.cell(row=3, column=last_col).font = Font(name="Calibri", size=8, italic=True, color=INPUT_BLUE)
        ws.cell(row=3, column=last_col).alignment = Alignment(horizontal="right", vertical="center")
    else:
        ws.merge_cells(f"B3:{last}3")
    ws["B3"].value = subtitle
    ws["B3"].font = Font(name="Calibri", size=10, color="475569")
    ws["B3"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 16
    apply_fill(ws, 4, 2, last_col, fill_teal)
    ws.row_dimensions[4].height = 3
    ws.row_dimensions[5].height = 8

# ═══════════════════════════════════════════════════
# SHEET 1: WRAP RATE CALCULATOR
# ═══════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "Wrap Rate Calculator"
ws1.sheet_view.showGridLines = False

ws1.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws1.page_setup.orientation = "portrait"
ws1.page_setup.paperSize = ws1.PAPERSIZE_A4
ws1.page_setup.fitToHeight = 1
ws1.page_setup.fitToWidth = 1
ws1.page_margins.left = 0.5
ws1.page_margins.right = 0.5
ws1.page_margins.top = 0.4
ws1.page_margins.bottom = 0.4

# Column widths: A=margin, B=labels, C=rates/values, D=formulas/notes, E=extra
widths = {'A': 2, 'B': 32, 'C': 16, 'D': 28, 'E': 16, 'F': 3}
for c, w in widths.items():
    ws1.column_dimensions[c].width = w

# ── Header ──
brand_header(ws1, "Wrap Rate Calculator",
             "Defense contractor rate build-up: from base pay to the hourly rate the government pays",
             5)

# ── INPUT SECTION ──
r = 6
ws1.merge_cells(f"B{r}:E{r}")
ws1.cell(row=r, column=2).value = "INPUT ASSUMPTIONS"
ws1.cell(row=r, column=2).font = font_section
ws1.cell(row=r, column=2).fill = fill_teal
ws1.cell(row=r, column=2).alignment = ac
apply_fill(ws1, r, 2, 5, fill_teal)
ws1.row_dimensions[r].height = 22

r = 7
# Headers for input section
for col, hdr in [(2, "Input"), (3, "Value"), (4, "Description")]:
    cell = ws1.cell(row=r, column=col)
    cell.value = hdr
    cell.font = Font(name="Calibri", size=9, bold=True, color=TEAL)
    cell.alignment = ac
    cell.border = bottom_thin
ws1.row_dimensions[r].height = 18

# Input rows with example values
inputs = [
    ("Base Hourly Rate", 65, "$#,##0.00", "Direct labor hourly rate: what the employee is paid per hour"),
    ("Fringe Benefit Rate %", 0.30, "0.0%", "FICA, health insurance, 401k, PTO, workers comp"),
    ("Overhead Rate %", 0.45, "0.0%", "Indirect costs: mgmt, facilities, IT, HR (applied to DL + Fringe)"),
    ("G&A Rate %", 0.12, "0.0%", "General & Administrative: exec, legal, finance (applied to Total Cost Input)"),
    ("Fee / Profit %", 0.08, "0.0%", "Contractor profit margin (negotiated; varies by contract type)"),
]

r = 8
input_rows = {}
for label, value, fmt, desc in inputs:
    ws1.cell(row=r, column=2).value = label
    ws1.cell(row=r, column=2).font = font_body
    ws1.cell(row=r, column=2).alignment = al_indent
    
    cell = ws1.cell(row=r, column=3)
    cell.value = value
    cell.font = font_input  # Blue text = input
    cell.number_format = fmt
    cell.alignment = ac
    cell.fill = fill_input
    cell.border = input_border
    
    ws1.cell(row=r, column=4).value = desc
    ws1.cell(row=r, column=4).font = font_note
    ws1.cell(row=r, column=4).alignment = al
    
    input_rows[label] = r
    ws1.row_dimensions[r].height = 22
    r += 1

# Cell references for formulas
base_cell = f"C{input_rows['Base Hourly Rate']}"
fringe_pct = f"C{input_rows['Fringe Benefit Rate %']}"
oh_pct = f"C{input_rows['Overhead Rate %']}"
ga_pct = f"C{input_rows['G&A Rate %']}"
fee_pct = f"C{input_rows['Fee / Profit %']}"

# ── OUTPUT / BUILDUP SECTION ──
r += 1
ws1.merge_cells(f"B{r}:E{r}")
ws1.cell(row=r, column=2).value = "RATE BUILD-UP (all formulas: change the inputs above to recalculate)"
ws1.cell(row=r, column=2).font = font_section
ws1.cell(row=r, column=2).fill = fill_navy
ws1.cell(row=r, column=2).alignment = ac
apply_fill(ws1, r, 2, 5, fill_navy)
ws1.row_dimensions[r].height = 22
r += 1

# Column headers
for col, hdr in [(2, "Component"), (3, "Amount"), (4, "Formula"), (5, "Running Total")]:
    cell = ws1.cell(row=r, column=col)
    cell.value = hdr
    cell.font = Font(name="Calibri", size=9, bold=True, color=TEAL)
    cell.alignment = ac
    cell.border = bottom_thin
ws1.row_dimensions[r].height = 18
r += 1

# Buildup rows
buildup = [
    ("Direct Labor (DL)", f"={base_cell}", "$#,##0.00",
     f'="$"&TEXT({base_cell},"#,##0.00")&"/hr (base rate)"',
     f"={base_cell}"),
    ("+ Fringe Benefits", f"={base_cell}*{fringe_pct}", "$#,##0.00",
     f'="DL × "&TEXT({fringe_pct},"0.0%")',
     f"={base_cell}+{base_cell}*{fringe_pct}"),
    ("\u2192 Fringe-Burdened Rate", f"={base_cell}*(1+{fringe_pct})", "$#,##0.00",
     '"DL + Fringe"',
     f"={base_cell}*(1+{fringe_pct})"),
    ("+ Overhead", f"={base_cell}*(1+{fringe_pct})*{oh_pct}", "$#,##0.00",
     f'="Fringe-Burdened × "&TEXT({oh_pct},"0.0%")',
     f"={base_cell}*(1+{fringe_pct})*(1+{oh_pct})"),
    ("\u2192 Total Cost Input (TCI)", f"={base_cell}*(1+{fringe_pct})*(1+{oh_pct})", "$#,##0.00",
     '"DL + Fringe + Overhead"',
     f"={base_cell}*(1+{fringe_pct})*(1+{oh_pct})"),
    ("+ G&A", f"={base_cell}*(1+{fringe_pct})*(1+{oh_pct})*{ga_pct}", "$#,##0.00",
     f'="TCI × "&TEXT({ga_pct},"0.0%")',
     f"={base_cell}*(1+{fringe_pct})*(1+{oh_pct})*(1+{ga_pct})"),
    ("\u2192 Total Cost", f"={base_cell}*(1+{fringe_pct})*(1+{oh_pct})*(1+{ga_pct})", "$#,##0.00",
     '"TCI + G&A"',
     f"={base_cell}*(1+{fringe_pct})*(1+{oh_pct})*(1+{ga_pct})"),
    ("+ Fee / Profit", f"={base_cell}*(1+{fringe_pct})*(1+{oh_pct})*(1+{ga_pct})*{fee_pct}", "$#,##0.00",
     f'="Total Cost × "&TEXT({fee_pct},"0.0%")',
     f"={base_cell}*(1+{fringe_pct})*(1+{oh_pct})*(1+{ga_pct})*(1+{fee_pct})"),
]

for label, formula, fmt, formula_desc, running in buildup:
    is_subtotal = label.startswith("\u2192")
    is_final = "Total Cost" in label and "Input" not in label and "+" not in label
    
    ws1.cell(row=r, column=2).value = label
    ws1.cell(row=r, column=2).font = font_body_bold if is_subtotal else font_body
    ws1.cell(row=r, column=2).alignment = al_indent
    
    cell_c = ws1.cell(row=r, column=3)
    cell_c.value = formula
    cell_c.font = font_formula_bold if is_subtotal else font_formula
    cell_c.number_format = fmt
    cell_c.alignment = ac
    cell_c.fill = fill_output if is_subtotal else fill_white
    
    ws1.cell(row=r, column=4).value = f"={formula_desc}" if not formula_desc.startswith("=") else formula_desc
    ws1.cell(row=r, column=4).font = font_note
    ws1.cell(row=r, column=4).alignment = al
    
    cell_e = ws1.cell(row=r, column=5)
    cell_e.value = running
    cell_e.font = font_formula_bold if is_subtotal else font_formula
    cell_e.number_format = fmt
    cell_e.alignment = ac
    cell_e.fill = fill_output if is_subtotal else fill_white
    
    if is_subtotal:
        for c in range(2, 6):
            ws1.cell(row=r, column=c).border = Border(
                top=Side(style="thin", color=BORDER_GRAY),
                bottom=Side(style="thin", color=BORDER_GRAY),
            )
    
    ws1.row_dimensions[r].height = 22
    r += 1

# ── FINAL RESULTS ──
r += 1
ws1.merge_cells(f"B{r}:E{r}")
ws1.cell(row=r, column=2).value = "FINAL RESULTS"
ws1.cell(row=r, column=2).font = font_section
ws1.cell(row=r, column=2).fill = fill_teal
ws1.cell(row=r, column=2).alignment = ac
apply_fill(ws1, r, 2, 5, fill_teal)
ws1.row_dimensions[r].height = 22
r += 1

# Fully Loaded Billing Rate
ws1.cell(row=r, column=2).value = "Fully Loaded Billing Rate"
ws1.cell(row=r, column=2).font = font_label_bold
ws1.cell(row=r, column=2).alignment = al_indent

cell = ws1.cell(row=r, column=3)
cell.value = f"={base_cell}*(1+{fringe_pct})*(1+{oh_pct})*(1+{ga_pct})*(1+{fee_pct})"
cell.font = font_result
cell.number_format = "$#,##0.00"
cell.alignment = ac
cell.fill = fill_result
cell.border = Border(
    left=Side(style="medium", color=TEAL),
    right=Side(style="medium", color=TEAL),
    top=Side(style="medium", color=TEAL),
    bottom=Side(style="medium", color=TEAL),
)

ws1.merge_cells(f"D{r}:E{r}")
ws1.cell(row=r, column=4).value = "Total Cost + Fee = what the government pays per hour"
ws1.cell(row=r, column=4).font = font_note
ws1.cell(row=r, column=4).alignment = al
ws1.row_dimensions[r].height = 28
r += 1

# Wrap Rate Multiplier
ws1.cell(row=r, column=2).value = "Wrap Rate Multiplier"
ws1.cell(row=r, column=2).font = font_label_bold
ws1.cell(row=r, column=2).alignment = al_indent

cell = ws1.cell(row=r, column=3)
cell.value = f"=(1+{fringe_pct})*(1+{oh_pct})*(1+{ga_pct})*(1+{fee_pct})"
cell.font = font_result
cell.number_format = "0.00x"
cell.alignment = ac
cell.fill = fill_result
cell.border = Border(
    left=Side(style="medium", color=TEAL),
    right=Side(style="medium", color=TEAL),
    top=Side(style="medium", color=TEAL),
    bottom=Side(style="medium", color=TEAL),
)

ws1.merge_cells(f"D{r}:E{r}")
ws1.cell(row=r, column=4).value = "Fully loaded rate ÷ base rate: how many times base pay the government pays"
ws1.cell(row=r, column=4).font = font_note
ws1.cell(row=r, column=4).alignment = al
ws1.row_dimensions[r].height = 28
r += 2

# Color legend
ws1.merge_cells(f"B{r}:E{r}")
ws1.cell(row=r, column=2).value = "★ Convention: Blue text = hardcoded inputs (change these)  |  Black text = formulas (auto-calculated)  |  Yellow background = input cells  |  Green = subtotals"
ws1.cell(row=r, column=2).font = font_note
ws1.cell(row=r, column=2).alignment = al
ws1.row_dimensions[r].height = 16
r += 1

ws1.merge_cells(f"B{r}:E{r}")
ws1.cell(row=r, column=2).value = FOOTER
ws1.cell(row=r, column=2).font = font_footer
ws1.cell(row=r, column=2).alignment = ac
ws1.print_area = f"A1:F{r}"

# ═══════════════════════════════════════════════════
# SHEET 2: RATE COMPARISON
# ═══════════════════════════════════════════════════
ws2 = wb.create_sheet("Rate Comparison")
ws2.sheet_view.showGridLines = False
ws2.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws2.page_setup.orientation = "landscape"
ws2.page_setup.paperSize = ws2.PAPERSIZE_A4
ws2.page_setup.fitToHeight = 1
ws2.page_setup.fitToWidth = 1
ws2.page_margins.left = 0.4
ws2.page_margins.right = 0.4
ws2.page_margins.top = 0.4
ws2.page_margins.bottom = 0.4

widths2 = {'A': 2, 'B': 34, 'C': 11, 'D': 11, 'E': 11, 'F': 12, 'G': 13, 'H': 46}
for c, w in widths2.items():
    ws2.column_dimensions[c].width = w

# Header
brand_header(ws2, "Wrap Rate Comparison by Contractor Type",
             "Typical ranges by type of business unit, built in the calculator's order so any row can be typed into the Wrap Rate Calculator tab",
             8)

# Table headers
r = 6
comp_headers = ["Contractor Type", "Typical Fringe", "Typical OH", "Typical G&A", "Typical Fee", "Typical\nWrap Rate", "Notes"]
comp_cols = list(range(2, 9))
for i, (hdr, col) in enumerate(zip(comp_headers, comp_cols)):
    cell = ws2.cell(row=r, column=col)
    cell.value = hdr
    cell.font = font_header
    cell.fill = fill_teal
    cell.alignment = ac
    cell.border = thin
ws2.row_dimensions[r].height = 28

# Data. Ranges are (low, high) percent. The wrap column is computed from them,
# in the calculator's order, so a row can never contradict itself again.
comp_rows = [
    ("Large Prime: services business unit\n(mission support, often on government sites)", (28, 34), (10, 30), (8, 12), (6, 9),
     "Staff often sit on government sites, so overhead is light. The same parent company can run a systems division at nearly double this multiplier."),
    ("Large Prime: engineering or systems unit\n(development, integration, C2, sensors)", (30, 38), (50, 80), (10, 14), (8, 10),
     "Engineering overhead carries labs, facilities, tools and IRAD at contractor sites. Each business unit negotiates its own forward pricing rates."),
    ("Services Integrator\n(large professional services firms)", (28, 35), (20, 40), (8, 12), (6, 9),
     "Scale spreads indirect costs thinly. Most quote separate on-site and off-site rates; on-site is lower."),
    ("Mid-Tier Contractor\n($100M to $1B, often grown out of set-asides)", (26, 34), (25, 45), (10, 15), (7, 10),
     "Less scale than an integrator, so G&A runs higher. The squeeze point when a firm outgrows small-business status."),
    ("Small Business\n(8(a), SDVOSB, HUBZone, WOSB)", (20, 30), (15, 40), (8, 18), (7, 12),
     "Lean overhead, but a small base can push G&A up. Many use one combined indirect rate. May lack DCAA-reviewed rates at first."),
    ("Cleared Specialty\n(cyber, SIGINT, cleared engineers)", (30, 38), (30, 55), (10, 14), (8, 12),
     "The talent premium shows up mostly in the base salary, not the multiplier. Clearance processing, SCIF space and retention sit in overhead."),
    ("OCONUS / Deployed\n(overseas support)", (35, 45), (55, 80), (12, 18), (8, 12),
     "Hazard, danger and post pay, housing, security and R&R travel. Much of it is priced as separate direct costs, so compare total cost per person."),
    ("FFRDC / UARC\n(nonprofit research centers)", (30, 40), (50, 75), (10, 16), (0, 3),
     "Nonprofit, so fee is small and based on need (DFARS 215.404-75). High overhead covers labs and research infrastructure."),
]


def _wrap(f, o, g, e):
    return (1 + f / 100) * (1 + o / 100) * (1 + g / 100) * (1 + e / 100)


def _pct(rng):
    return f"{rng[0]}–{rng[1]}%"


comp_data = []
for ctype, fr, oh, ga, fee, notes in comp_rows:
    lo = _wrap(fr[0], oh[0], ga[0], fee[0])
    hi = _wrap(fr[1], oh[1], ga[1], fee[1])
    wrap = f"{round(lo + 1e-9, 1):.1f}x – {round(hi + 1e-9, 1):.1f}x"
    fee_txt = _pct(fee) + (" (need-based)" if fee[0] == 0 else "")
    comp_data.append((ctype, _pct(fr), _pct(oh), _pct(ga), fee_txt, wrap, notes))

r = 7
for i, (ctype, fringe, oh, ga, fee, wrap, notes) in enumerate(comp_data):
    alt = fill_light_teal if i % 2 == 0 else fill_white
    vals = [ctype, fringe, oh, ga, fee, wrap, notes]
    for j, (val, col) in enumerate(zip(vals, comp_cols)):
        cell = ws2.cell(row=r, column=col)
        cell.value = val
        cell.border = thin
        cell.fill = alt
        if j == 0:
            cell.font = font_body_bold
            cell.alignment = al_indent
        elif j == 6:
            cell.font = Font(name="Calibri", size=8, color=MID_GRAY)
            cell.alignment = al_top
        else:
            cell.font = font_body_sm
            cell.alignment = ac
    ws2.row_dimensions[r].height = 42
    r += 1

# Reading notes
ws2_notes = [
    ("HOW TO READ THIS TABLE",
     "Rates build up in the calculator's order: fringe on direct labor, overhead on labor plus fringe, G&A on total cost input, fee on total cost. "
     "The wrap rate column multiplies the four ranges together at their low and high ends (rounded), and includes fee. Many companies quote a wrap without fee, "
     "fold fringe into overhead, or use one combined indirect rate, so compare structures before comparing numbers."),
    ("ONE COMPANY, MANY WRAP RATES",
     "A large prime is not one rate. Each business unit and site has its own overhead pools and its own forward pricing rate agreement, and corporate G&A "
     "is allocated to each (CAS 403, CAS 410). At one large prime, a services business unit ran about 1.7x while a C2 systems division under the same "
     "parent ran about 2.9x. Always ask which business unit and which site a rate comes from."),
    ("ON-SITE VS OFF-SITE",
     "Work on a government site usually carries a lower overhead rate, because the government supplies the desk, the building and the utilities. "
     "Expect a services contractor to price both. Factory-floor production labor usually wraps higher still, because manufacturing overhead carries the plant and equipment."),
    ("VERIFY",
     "These are typical ranges, not any company's rates. DCAA advises against comparing indirect rates between organizations at the rate level. "
     "For a real proposal, check the contractor's current FPRA or FPRR."),
]
r += 1
for head, text in ws2_notes:
    ws2.cell(row=r, column=2).value = head
    ws2.cell(row=r, column=2).font = Font(name="Calibri", size=8, bold=True, color=TEAL)
    ws2.cell(row=r, column=2).alignment = al_top
    ws2.merge_cells(f"C{r}:H{r}")
    ws2.cell(row=r, column=3).value = text
    ws2.cell(row=r, column=3).font = Font(name="Calibri", size=8, color=DARK_GRAY)
    ws2.cell(row=r, column=3).alignment = al_top
    ws2.row_dimensions[r].height = 34
    r += 1

r += 1
ws2.merge_cells(f"B{r}:H{r}")
ws2.cell(row=r, column=2).value = FOOTER
ws2.cell(row=r, column=2).font = font_footer
ws2.cell(row=r, column=2).alignment = ac
ws2.print_area = f"A1:H{r}"

# ═══════════════════════════════════════════════════
# SHEET 3: INSTRUCTIONS
# ═══════════════════════════════════════════════════
ws3 = wb.create_sheet("Instructions")
ws3.sheet_view.showGridLines = False
ws3.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws3.page_setup.orientation = "portrait"
ws3.page_setup.paperSize = ws3.PAPERSIZE_A4
ws3.page_setup.fitToHeight = 1
ws3.page_setup.fitToWidth = 1

widths3 = {'A': 2, 'B': 4, 'C': 80, 'D': 2}
for c, w in widths3.items():
    ws3.column_dimensions[c].width = w

# Header
brand_header(ws3, "Understanding Wrap Rates in Defense Contracting",
             "What the multiplier means, how the government checks it, and how to read it in a source selection", 3)

# Content sections
sections = [
    ("WHAT WRAP RATES ARE & WHY THEY MATTER", [
        "A wrap rate (or burden rate/multiplier) is the total cost the government pays per hour of labor divided by the employee's base hourly wage. It captures all indirect costs, overhead, and profit layered on top of direct labor.",
        "",
        "Why it matters to acquisition professionals:",
        "  ▸ It determines the true cost of every labor hour on your contract",
        "  ▸ A $65/hr engineer may cost the government anywhere from about $105 to $200/hr fully loaded, depending on the business unit and work site",
        "  ▸ Understanding wrap rates is essential for cost realism analysis in source selection",
        "  ▸ Wrap rate differences between competitors explain most of the price spread in professional services",
        "  ▸ Unrealistically low wrap rates signal an unrealistic proposal (buying in)",
    ]),
    ("HOW THE GOVERNMENT CHECKS RATES", [
        "DCAA audits contractor rates and DCMA's administrative contracting officer (ACO) negotiates them. The main mechanisms:",
        "",
        "  ▸ FPRA (Forward Pricing Rate Agreement): a written agreement between the contractor and the ACO on the indirect rates to use in pricing proposals for a future period. The gold standard.",
        "  ▸ FPRR (Forward Pricing Rate Recommendation): when there is no FPRA, the ACO recommends rates for negotiators to use, informed by DCAA's audit. Less binding than an FPRA.",
        "  ▸ Incurred cost audits: DCAA compares each year's actual indirect rates with the provisional billing rates. The result sets final rates and any adjustments.",
        "  ▸ CAS compliance: the Cost Accounting Standards (CAS 401 to 420) require consistency in how costs are estimated, accumulated and reported.",
        "",
        "Key point: if a contractor has no audited rate structure, ask for its rate build-up with supporting documentation. Provisional rates may apply.",
    ]),
    ("WHAT 'UNCOMPETITIVE WRAP RATES' MEANS IN SOURCE SELECTION", [
        "During source selection, evaluators performing cost realism analysis should examine wrap rates to ensure:",
        "",
        "  ▸ Rates are consistent with the contractor's FPRA/FPRR (if available)",
        "  ▸ Rates are within norms for the same kind of business unit and the same work site (on-site or off-site), not just the same company size",
        "  ▸ Rates adequately cover the contractor's real cost of doing business",
        "  ▸ Unusually LOW rates may indicate buy-in pricing (contractor plans to raise rates later)",
        "  ▸ Unusually HIGH rates may indicate over-staffing or inefficiency",
        "",
        "Red flags in wrap rate analysis:",
        "  ▸ A wrap rate below about 1.5x on off-site work, fee included (hard to sustain)",
        "  ▸ Overhead rate falling year over year with no explanation",
        "  ▸ Fringe far below the company's own history (benefit cuts can mean turnover)",
        "  ▸ G&A above 20% (worth a question about the cost structure)",
    ]),
    ("HOW TO CHECK IF A COMPETITOR'S PRICING IS OVER-LEVELED", [
        "Over-leveling is proposing senior labor for work junior staff could do. It is a common pricing tactic:",
        "",
        "  ▸ Compare proposed labor categories against the PWS/SOW task complexity",
        "  ▸ Check if the proposed labor mix matches industry norms for similar work",
        "  ▸ Verify years of experience claims are realistic for proposed categories",
        "  ▸ Compare against GSA Schedule rates for equivalent labor categories",
        "  ▸ Calculate effective hourly rates and compare against BLS data for the region",
        "",
        "Tip: Request detailed labor category descriptions with minimum qualifications. Map each category to specific PWS tasks. If Level IV engineers are proposed for Level II work, that's over-leveling.",
    ]),
    ("KEY REGULATORY REFERENCES", [
        "  ▸ FAR 31.201: composition of total cost",
        "  ▸ FAR 31.203: indirect costs and how to allocate them",
        "  ▸ FAR 31.205: selected costs, allowable and unallowable",
        "  ▸ FAR 15.404-1(d): cost realism analysis",
        "  ▸ CAS 401: consistency in estimating, accumulating and reporting costs",
        "  ▸ CAS 402: consistency in allocating costs incurred for the same purpose",
        "  ▸ CAS 403: allocating home office (corporate) expenses to business units",
        "  ▸ CAS 410: allocating business unit G&A",
        "  ▸ CAS 418: allocating direct and indirect costs",
        "  ▸ DFARS 215.404-71: weighted guidelines for profit and fee",
        "  ▸ DCAA Contract Audit Manual: chapter 6 (incurred costs), chapter 9 (cost estimates and price proposals)",
    ]),
]

r = 6
for title, lines in sections:
    ws3.merge_cells(f"B{r}:C{r}")
    ws3.cell(row=r, column=2).value = title
    ws3.cell(row=r, column=2).font = font_section
    ws3.cell(row=r, column=2).fill = fill_teal
    ws3.cell(row=r, column=2).alignment = ac
    apply_fill(ws3, r, 2, 3, fill_teal)
    ws3.row_dimensions[r].height = 22
    r += 1
    
    for line in lines:
        ws3.merge_cells(f"B{r}:C{r}")
        cell = ws3.cell(row=r, column=2)
        if line == "":
            ws3.row_dimensions[r].height = 6
        elif line.startswith("  ▸"):
            cell.value = line
            cell.font = font_inst_body
            cell.alignment = al_top
            ws3.row_dimensions[r].height = 14 * (1 + (len(line) - 1) // 88) + 2
        elif line.endswith(":"):
            cell.value = line
            cell.font = font_inst_bold
            cell.alignment = al_top
            ws3.row_dimensions[r].height = 14 * (1 + (len(line) - 1) // 84) + 3
        else:
            cell.value = line
            cell.font = font_inst_body
            cell.alignment = al_top
            ws3.row_dimensions[r].height = 14 * (1 + (len(line) - 1) // 88) + 2
        r += 1
    r += 1  # Gap between sections

# Footer
ws3.merge_cells(f"B{r}:C{r}")
ws3.cell(row=r, column=2).value = FOOTER
ws3.cell(row=r, column=2).font = font_footer
ws3.cell(row=r, column=2).alignment = ac

ws3.print_area = f"A1:D{r}"

# ── Save ──
for out_dir in (ROOT / "client/public/products/pack3-finance-cheat-sheets", ROOT / "products/pack3-finance-cheat-sheets"):
    output = out_dir / "wrap-rate-breakdown.xlsx"
    wb.save(output)
    print(f"Saved: {output}")
